# -*- coding: utf-8 -*-
from openpyxl import load_workbook
wb = load_workbook(r"C:\Users\USUARIO\Documents\deepseek-harness\default-workspace\Control-Compras\control-compras.xlsx", data_only=True)
ws = wb["Tickets"]
hoja = """<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Reporte</title>
<style>body{font-family:sans-serif;padding:20px;background:#f4f6f9;color:#1c2a38;}h1{color:#1f4e79;}table{border-collapse:collapse;width:100%;}th{background:#1f4e79;color:#fff;}td,th{padding:8px;border:1px solid #ccc;}</style>
</head><body>
<h1>Control de Compras — Reporte Estático</h1>
<h2>Predicción: Martes y Sábado</h2>
<p>Generado directamente desde control-compras.xlsx. No requiere servidor ni Python abierto.</p>
<h3>Tickets registrados</h3>
<table><tr><th>Comercio</th><th>Fecha</th><th>Total</th><th>Estado</th></tr>"""
for row in ws.iter_rows(min_row=2, max_row=min(15, ws.max_row), values_only=True):
    if row[0] is None: continue
    comercio = str(row[0] or "")
    fecha = str(row[1] or "")
    total = str(row[3] or "")
    estado = str(row[6] or "")
    hoja += f"<tr><td>{comercio}</td><td>{fecha}</td><td>{total}</td><td>{estado}</td></tr>\n"
hoja += "</table><p><strong>Para predicciones:</strong> revisá los últimos tickets cargados y estimá compras antes del martes y sábado según frecuencia.</p></body></html>"
with open(r"C:\Users\USUARIO\Documents\deepseek-harness\default-workspace\Control-Compras\reporte.html", "w", encoding="utf-8") as f:
    f.write(hoja)
print("reporte.html creado")
