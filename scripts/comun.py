# -*- coding: utf-8 -*-
"""Constantes y utilidades compartidas del sistema de control de compras.

Fuente unica de verdad para: rutas, columnas de cada hoja, limites de filas
y la clave que detecta tickets duplicados.
"""
import csv
import datetime as dt
import glob
import hashlib
import os
import shutil
import unicodedata
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, "control-compras.xlsx")
DIR_TICKETS = os.path.join(BASE, "tickets")
DIR_ENTRADA = os.path.join(DIR_TICKETS, "entrada")
DIR_ARCHIVADOS = os.path.join(DIR_TICKETS, "archivados")
DIR_LECTURA = os.path.join(BASE, "lecturas")
MANIFIESTO = os.path.join(DIR_TICKETS, "manifiesto.csv")
ANULADOS = os.path.join(DIR_TICKETS, "anulados.csv")
COPIAS_MAX = 20
MANIFIESTO_CAB = ["id", "fecha", "archivo", "bytes", "sha256", "registrado_en"]
ANULADOS_CAB = ["id", "fecha", "comercio", "total", "archivo", "motivo", "anulado_en"]

# Configuración de Tesseract OCR (para lectura automática de tickets de foto)
TESSERACT_CMD = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# --- motor del entorno para convertir PDF a imagen
NODE = r"C:\Users\USUARIO\.dsh\dsh-runtimes\dsh-primary-runtime\dependencies\node\bin\node.exe"
LIBREOFFICE_CLI = (r"C:\Users\USUARIO\AppData\Local\Programs\DeepSeek Harness\resources"
                   r"\app.asar.unpacked\dsh\node_modules\@deepseek-ai\libreoffice-kit"
                   r"\lib\cli.js")

EXT_IMAGEN = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff")
EXT_PDF = (".pdf",)

# --- capacidad de la plantilla
CAP_CAT = 24         # categorias admitidas en Listas / Resumen
CAP_PAGO = 14        # formas de pago
CAP_UNID = 10        # unidades
FILAS = 5000         # alcance de los SUMIFS
FILAS_FORM = 1000    # filas de Compras con formulas ya escritas
FILAS_TK = 400       # filas de Tickets con formulas ya escritas
DUP_HASTA = 1000     # alcance del control de duplicados

# --- formato
MONEDA = '"$"#,##0.00'
FECHA = 'DD/MM/YYYY'
PCT = '0.0%'

# --- columnas (el orden define la posicion real en la hoja)
COMPRAS_COLS = [
    "ID ticket", "Linea", "Fecha", "Mes", "Comercio", "Categoria", "Descripcion",
    "Cant.", "Unid.", "Precio unit.", "Descuento", "Importe", "IVA %", "IVA $",
    "Forma de pago", "Moneda", "N. ticket", "Archivo", "Notas",
]
TICKETS_COLS = [
    "ID ticket", "Fecha", "Comercio", "Total ticket", "Total registrado",
    "Diferencia", "Estado", "N. ticket", "Forma de pago", "Moneda",
    "Archivo escaneado", "Notas", "Duplicado",
]
LISTAS_COLS = ["Categorias", "Formas de pago", "Unidades", "Moneda"]

IC = {nombre: i for i, nombre in enumerate(COMPRAS_COLS)}        # 0-based
IT = {nombre: i for i, nombre in enumerate(TICKETS_COLS)}        # 0-based


def cl(campo):
    """Letra de columna de Compras."""
    return get_column_letter(IC[campo] + 1)


def tl(campo):
    """Letra de columna de Tickets."""
    return get_column_letter(IT[campo] + 1)


# --- listas por defecto
CATEGORIAS = [
    "Supermercado", "Almacen y bebidas", "Carniceria", "Verduleria y frutas",
    "Panaderia", "Restaurantes y bares", "Transporte y combustible",
    "Hogar y ferreteria", "Ropa y calzado", "Salud y farmacia", "Tecnologia",
    "Ocio y entretenimiento", "Servicios y suministros", "Educacion",
    "Mascotas", "Cuidado personal", "Regalos", "Impuestos y tasas",
    "Viajes", "Otros",
]
FORMAS_PAGO = [
    "Efectivo", "Tarjeta debito", "Tarjeta credito", "Transferencia",
    "Billetera virtual", "QR / Mercado Pago", "Cheque", "Debito automatico",
    "Cuenta corriente", "Crypto", "Vale / cupon", "Otro",
]
UNIDADES = ["ud", "kg", "g", "L", "ml", "m", "m2", "caja", "pack", "servicio"]
MONEDAS = ["ARS", "USD", "EUR", "BRL", "UYU", "CLP", "PYG", "BOB", "MXN", "COP", "PEN"]


def normalizar(texto):
    """Espacios colapsados y sin espacios sobrantes."""
    return " ".join(str(texto or "").split()).strip()


def slug(texto):
    """Texto apto para nombre de archivo: sin acentos ni simbolos."""
    sin_acentos = unicodedata.normalize("NFKD", str(texto or ""))
    sin_acentos = "".join(c for c in sin_acentos if not unicodedata.combining(c))
    limpio = "".join(c if c.isalnum() else "-" for c in sin_acentos.lower())
    return "-".join(p for p in limpio.split("-") if p) or "ticket"


def clave_duplicado(fecha, comercio, total, numero):
    """Clave de duplicado. Debe coincidir con el COUNTIFS de la hoja Tickets.

    La hoja compara fecha + comercio + total + n. de ticket (sin distinguir
    mayusculas). Aqui se replica esa misma comparacion.
    """
    if fecha in (None, ""):
        raise ValueError("fecha vacia")
    fecha = formatear_fecha(fecha)  # normaliza datetime -> date
    return (
        fecha,
        normalizar(comercio).lower(),
        round(float(total), 2),
        normalizar(numero).lower(),
    )


def nombre_archivo(fecha, comercio, total, extension):
    """Nombre normalizado del escaneo: AAAA-MM-DD_comercio_total.ext"""
    if isinstance(fecha, str):
        fecha = dt.date.fromisoformat(fecha)
    ext = extension if extension.startswith(".") else "." + extension
    return f"{fecha.isoformat()}_{slug(comercio)}_{float(total):.2f}{ext.lower()}"


def ruta_legible(ruta):
    """Ruta relativa a BASE si es posible, con barras normales."""
    try:
        rel = os.path.relpath(ruta, BASE)
    except ValueError:
        return ruta.replace("\\", "/")
    return rel.replace("\\", "/") if not rel.startswith("..") else ruta.replace("\\", "/")


def formatear_fecha(valor):
    """Acepta date/datetime o texto YYYY-MM-DD / DD/MM/YYYY."""
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    texto = normalizar(valor)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"fecha no reconocida: {valor!r} (usa AAAA-MM-DD)")


# ------------------------------------------------------------------ hojas

def leer_tabla(ws, indices, max_filas=FILAS):
    """Lee una hoja como lista de dicts (solo filas con la columna A ocupada).

    indices: mapa {campo: posicion 0-based} (los de COMPRAS_COLS/TICKETS_COLS).
    Anade '_fila' con el numero de fila real de la hoja.
    """
    ancho = max(indices.values()) + 1
    salida = []
    for r, fila in enumerate(ws.iter_rows(min_row=2, max_row=max_filas, min_col=1,
                                          max_col=ancho, values_only=True), start=2):
        if fila[0] in (None, ""):
            continue
        salida.append({"_fila": r, **{c: fila[i] for c, i in indices.items()}})
    return salida


def fila_libre(ws, desde=2, tope=FILAS):
    """Primera fila sin dato en la columna A, a partir de 'desde'."""
    for r in range(desde, tope + 1):
        if ws.cell(row=r, column=1).value in (None, ""):
            return r
    raise ValueError(f"la hoja {ws.title} esta llena (limite {tope} filas)")


def importe_linea(cantidad, precio, descuento):
    """Mismo calculo que la formula de la hoja Compras."""
    return round((cantidad or 0) * (precio or 0) - (descuento or 0), 2)


def iva_linea(importe, iva_pct):
    """IVA incluido en el importe, igual que la formula de la columna 'IVA $'."""
    if iva_pct in (None, "") or importe in (None, ""):
        return None
    return round(importe - importe / (1 + iva_pct / 100), 2)


# ------------------------------------------------------------------ archivos

def sha256(ruta, trozo=1024 * 1024):
    """Huella del archivo, para comprobar que el original no ha cambiado."""
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        while True:
            datos = f.read(trozo)
            if not datos:
                break
            h.update(datos)
    return h.hexdigest()


def copia_seguridad(ruta, motivo="auto", maximo=COPIAS_MAX):
    """Copia el libro antes de sobrescribirlo y rota las copias antiguas.

    Las copias viven en una carpeta 'copias' junto al libro.
    """
    if not os.path.exists(ruta):
        return None
    dir_copias = os.path.join(os.path.dirname(os.path.abspath(ruta)), "copias")
    os.makedirs(dir_copias, exist_ok=True)
    sello = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    base = os.path.splitext(os.path.basename(ruta))[0]
    destino = os.path.join(dir_copias, f"{base}_{sello}_{motivo}.xlsx")
    n = 2
    while os.path.exists(destino):
        destino = os.path.join(dir_copias, f"{base}_{sello}_{motivo}-{n}.xlsx")
        n += 1
    shutil.copy2(ruta, destino)
    for viejo in sorted(glob.glob(os.path.join(dir_copias, f"{base}_*.xlsx")))[:-maximo]:
        os.remove(viejo)
    return destino


# ------------------------------------------------------------------ manifiesto

def leer_manifiesto(ruta=MANIFIESTO):
    """Entradas del manifiesto de escaneos archivados."""
    if not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8", newline="") as f:
        return [dict(fila) for fila in csv.DictReader(f)]


def anotar_manifiesto(entradas, ruta=MANIFIESTO):
    """Anade entradas (id, fecha, archivo, bytes, sha256, registrado_en)."""
    if not entradas:
        return ruta
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    nuevo = not os.path.exists(ruta)
    with open(ruta, "a", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=MANIFIESTO_CAB)
        if nuevo:
            w.writeheader()
        for e in entradas:
            w.writerow({k: e.get(k, "") for k in MANIFIESTO_CAB})
    return ruta


def leer_anulados(ruta):
    """Tickets anulados: se conservan para no borrar la historia."""
    if not ruta or not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8", newline="") as f:
        return [dict(fila) for fila in csv.DictReader(f)]


def anotar_anulado(entrada, ruta):
    """Registra la anulacion de un ticket (el escaneo NO se borra)."""
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    nuevo = not os.path.exists(ruta)
    with open(ruta, "a", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=ANULADOS_CAB)
        if nuevo:
            w.writeheader()
        w.writerow({k: entrada.get(k, "") for k in ANULADOS_CAB})
    return ruta

