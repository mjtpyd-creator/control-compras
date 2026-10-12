# -*- coding: utf-8 -*-
"""Crea el libro de control de compras por tickets.

Genera control-compras.xlsx con las hojas:
  Resumen | Compras | Tickets | Presupuesto | Listas

Uso:
  python crear_libro.py            # se niega si ya hay tickets registrados
  python crear_libro.py --forzar   # recrea el libro VACIO (borra los datos)
"""
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import comun as co

XLSX = co.XLSX
FILAS = co.FILAS
FILAS_FORM = co.FILAS_FORM
FILAS_TK = co.FILAS_TK
DUP_HASTA = co.DUP_HASTA
CL, TL = co.cl, co.tl

# --- reparto de la hoja Resumen
R_TOT_TIT = 5
R_TOT_INI = 6
R_TOT_FIN = 12
R_MES_TIT = 14
R_MES_CAB = 15
R_MES_INI = 16
R_MES_FIN = R_MES_INI + 11
R_MES_TOT = R_MES_FIN + 1
R_CAT_TIT = R_MES_TOT + 2
R_CAT_CAB = R_CAT_TIT + 1
R_CAT_INI = R_CAT_CAB + 1
R_CAT_FIN = R_CAT_INI + co.CAP_CAT - 1
R_CAT_SIN = R_CAT_FIN + 1
R_CAT_TOT = R_CAT_SIN + 1
R_PAG_TIT = R_CAT_TOT + 2
R_PAG_CAB = R_PAG_TIT + 1
R_PAG_INI = R_PAG_CAB + 1
R_PAG_FIN = R_PAG_INI + co.CAP_PAGO - 1
R_PAG_TOT = R_PAG_FIN + 1

AZUL = "1F4E79"
AZUL_CLARO = "DDEBF7"
ROJO = "FFC7CE"
ROJO_TXT = "9C0006"
VERDE = "C6EFCE"
VERDE_TXT = "006100"
ENTRADA = "FFF2CC"

f_titulo = Font(bold=True, size=14, color=AZUL)
f_cab = Font(bold=True, color="FFFFFF", size=11)
f_seccion = Font(bold=True, color="FFFFFF", size=11)
f_bold = Font(bold=True)
f_nota = Font(italic=True, size=9, color="808080")

fill_cab = PatternFill("solid", fgColor=AZUL)
fill_total = PatternFill("solid", fgColor=AZUL_CLARO)
fill_entrada = PatternFill("solid", fgColor=ENTRADA)
fill_rojo = PatternFill("solid", fgColor=ROJO)
fill_verde = PatternFill("solid", fgColor=VERDE)

thin = Side(style="thin", color="BFBFBF")
borde = Border(left=thin, right=thin, top=thin, bottom=thin)

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
ANCHOS_COMPRAS = [11, 7, 11, 9, 22, 22, 34, 8, 8, 12, 11, 12, 8, 11, 16, 9, 14, 26, 30]
ANCHOS_TICKETS = [11, 11, 24, 14, 16, 12, 12, 15, 16, 9, 30, 30, 14]


def cabecera(ws, fila, etiquetas, anchos=None):
    for i, texto in enumerate(etiquetas, start=1):
        c = ws.cell(row=fila, column=i, value=texto)
        c.font = f_cab
        c.fill = fill_cab
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borde
    ws.row_dimensions[fila].height = 30
    if anchos:
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho


def titulo_seccion(ws, celda, texto):
    c = ws[celda]
    c.value = texto
    c.font = f_seccion
    c.fill = fill_cab


def fila_etiqueta(ws, fila, etiqueta, formula=None, formato=None, negrita=True):
    c = ws.cell(row=fila, column=1, value=etiqueta)
    if negrita:
        c.font = f_bold
    if formula is not None:
        v = ws.cell(row=fila, column=2, value=formula)
        v.fill = fill_total
        if formato:
            v.number_format = formato


def tiene_datos(ruta):
    """True si el libro existente ya tiene tickets registrados."""
    if not os.path.exists(ruta):
        return False
    try:
        wb = load_workbook(ruta, read_only=True)
    except Exception:
        return False
    try:
        # solo columnas de entrada: a partir de ahi hay formulas de plantilla
        for hoja, limite in (("Compras", 3), ("Tickets", 4)):
            if hoja in wb.sheetnames:
                ws = wb[hoja]
                for fila in ws.iter_rows(min_row=2, max_row=60, min_col=1, max_col=limite):
                    if any(c.value not in (None, "") for c in fila):
                        return True
    finally:
        wb.close()
    return False


if tiene_datos(XLSX) and "--forzar" not in sys.argv:
    sys.exit(
        "ABORTADO: " + XLSX + " ya contiene tickets registrados.\n"
        "Este script recrea el libro VACIO y perderias los datos.\n"
        "Si de verdad quieres vaciarlo: python crear_libro.py --forzar\n"
        "(haz antes una copia de seguridad del archivo)."
    )

wb = Workbook()

# ---------------------------------------------------------------- Listas
ws_listas = wb.active
ws_listas.title = "Listas"
cabecera(ws_listas, 1, co.LISTAS_COLS, [30, 24, 14, 10])
for col, valores in zip(range(1, 5),
                        [co.CATEGORIAS, co.FORMAS_PAGO, co.UNIDADES, co.MONEDAS]):
    for i, v in enumerate(valores, start=2):
        ws_listas.cell(row=i, column=col, value=v)
ws_listas["F1"] = (f"Listas editables. Capacidad maxima: {co.CAP_CAT} categorias, "
                   f"{co.CAP_PAGO} formas de pago, {co.CAP_UNID} unidades. "
                   "Si anades una categoria nueva, aparece sola en el Resumen y el Presupuesto.")
ws_listas["F1"].font = f_nota
ws_listas.column_dimensions["F"].width = 95

# ---------------------------------------------------------------- Compras
ws = wb.create_sheet("Compras")
cabecera(ws, 1, co.COMPRAS_COLS, ANCHOS_COMPRAS)
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(co.COMPRAS_COLS))}1"

for r in range(2, FILAS_FORM + 1):
    ws.cell(row=r, column=co.IC["Mes"] + 1,
            value=f'=IF(${CL("Fecha")}{r}="","",YEAR(${CL("Fecha")}{r})&"-"&TEXT(MONTH(${CL("Fecha")}{r}),"00"))')
    ws.cell(row=r, column=co.IC["Importe"] + 1,
            value=(f'=IF(${CL("Cant.")}{r}="","",'
                   f'ROUND(${CL("Cant.")}{r}*${CL("Precio unit.")}{r}-N(${CL("Descuento")}{r}),2))'))
    ws.cell(row=r, column=co.IC["IVA $"] + 1,
            value=(f'=IF(OR(${CL("Importe")}{r}="",${CL("IVA %")}{r}=""),"",'
                   f'ROUND(${CL("Importe")}{r}-${CL("Importe")}{r}/(1+${CL("IVA %")}{r}/100),2))'))
    ws.cell(row=r, column=co.IC["Fecha"] + 1).number_format = co.FECHA
    for campo in ("Cant.", "Precio unit.", "Descuento", "Importe", "IVA $"):
        ws.cell(row=r, column=co.IC[campo] + 1).number_format = co.MONEDA
    ws.cell(row=r, column=co.IC["IVA %"] + 1).number_format = '0.0"%"'
    for col in range(1, len(co.COMPRAS_COLS) + 1):
        ws.cell(row=r, column=col).border = borde

dv_rangos = [
    (DataValidation(type="list", formula1=f"Listas!$A$2:$A${co.CAP_CAT + 1}",
                    allow_blank=True, showDropDown=False), CL("Categoria")),
    (DataValidation(type="list", formula1=f"Listas!$B$2:$B${co.CAP_PAGO + 1}",
                    allow_blank=True, showDropDown=False), CL("Forma de pago")),
    (DataValidation(type="list", formula1=f"Listas!$C$2:$C${co.CAP_UNID + 1}",
                    allow_blank=True, showDropDown=False), CL("Unid.")),
    (DataValidation(type="list", formula1="Listas!$D$2:$D$12",
                    allow_blank=True, showDropDown=False), CL("Moneda")),
]
for dv, letra in dv_rangos:
    ws.add_data_validation(dv)
    dv.add(f"{letra}2:{letra}5000")

ws["U1"] = ("Escribe o pega aqui una fila por producto del ticket. "
            "'Mes', 'Importe' e 'IVA $' se calculan solos.")
ws["U1"].font = f_nota

# ---------------------------------------------------------------- Tickets
wt = wb.create_sheet("Tickets")
cabecera(wt, 1, co.TICKETS_COLS, ANCHOS_TICKETS)
wt.freeze_panes = "A2"
wt.auto_filter.ref = f"A1:{get_column_letter(len(co.TICKETS_COLS))}1"

CMP = "Compras"
for r in range(2, FILAS_TK + 1):
    wt.cell(row=r, column=co.IT["Total registrado"] + 1, value=(
        f'=IF(${TL("ID ticket")}{r}="","",'
        f'ROUND(SUMIFS({CMP}!${CL("Importe")}$2:${CL("Importe")}${FILAS},'
        f'{CMP}!${CL("ID ticket")}$2:${CL("ID ticket")}${FILAS},'
        f'${TL("ID ticket")}{r}),2))'))
    wt.cell(row=r, column=co.IT["Diferencia"] + 1, value=(
        f'=IF(${TL("ID ticket")}{r}="","",'
        f'ROUND(${TL("Total ticket")}{r}-${TL("Total registrado")}{r},2))'))
    wt.cell(row=r, column=co.IT["Estado"] + 1, value=(
        f'=IF(${TL("ID ticket")}{r}="","",'
        f'IF(ABS(${TL("Diferencia")}{r})<=0.02,"OK","REVISAR"))'))
    wt.cell(row=r, column=co.IT["Duplicado"] + 1, value=(
        f'=IF(${TL("ID ticket")}{r}="","",'
        f'IF(COUNTIFS(${TL("Fecha")}$2:${TL("Fecha")}${DUP_HASTA},${TL("Fecha")}{r},'
        f'${TL("Comercio")}$2:${TL("Comercio")}${DUP_HASTA},${TL("Comercio")}{r},'
        f'${TL("Total ticket")}$2:${TL("Total ticket")}${DUP_HASTA},${TL("Total ticket")}{r},'
        f'${TL("N. ticket")}$2:${TL("N. ticket")}${DUP_HASTA},${TL("N. ticket")}{r})>1,'
        f'"DUPLICADO",""))'))
    wt.cell(row=r, column=co.IT["Fecha"] + 1).number_format = co.FECHA
    for campo in ("Total ticket", "Total registrado", "Diferencia"):
        wt.cell(row=r, column=co.IT[campo] + 1).number_format = co.MONEDA
    for col in range(1, len(co.TICKETS_COLS) + 1):
        wt.cell(row=r, column=col).border = borde

wt.conditional_formatting.add(
    f"A2:{get_column_letter(len(co.TICKETS_COLS))}{FILAS_TK}",
    FormulaRule(formula=[f'${TL("Duplicado")}2="DUPLICADO"'], fill=fill_rojo,
                font=Font(color=ROJO_TXT, bold=True)))
wt.conditional_formatting.add(
    f"A2:{get_column_letter(len(co.TICKETS_COLS))}{FILAS_TK}",
    FormulaRule(formula=[f'AND($A2<>"",${TL("Diferencia")}2<>0)'], fill=fill_rojo,
                font=Font(color=ROJO_TXT)))
wt.conditional_formatting.add(
    f"A2:{get_column_letter(len(co.TICKETS_COLS))}{FILAS_TK}",
    FormulaRule(formula=[f'${TL("Estado")}2="OK"'], fill=fill_verde, font=Font(color=VERDE_TXT)))
wt["O1"] = ("Una fila por ticket. 'Total ticket' = lo que dice el papel; 'Total registrado' = "
            "suma de sus lineas en Compras. Verde = cuadra, rojo = REVISAR o DUPLICADO.")
wt["O1"].font = f_nota
wt.column_dimensions["O"].width = 70

# ---------------------------------------------------------------- Resumen
wr = wb.create_sheet("Resumen", 0)
wr.column_dimensions["A"].width = 34
for col in "BCD":
    wr.column_dimensions[col].width = 17
wr.column_dimensions["E"].width = 14

wr["A1"] = "CONTROL DE COMPRAS POR TICKETS"
wr["A1"].font = f_titulo
wr["A2"] = "Ejercicio analizado:"
wr["A2"].font = f_bold
wr["B2"] = "=YEAR(TODAY())"
wr["B2"].font = f_bold
wr["B2"].fill = fill_entrada
wr["B2"].number_format = "0"
wr["A3"] = "Moneda:"
wr["A3"].font = f_bold
wr["B3"] = "ARS"
wr["B3"].fill = fill_entrada

ANIO = 'DATE($B$2,1,1)'
FIN_ANIO = 'DATE($B$2,12,31)'
RANGO_C = f"Compras!${CL('Fecha')}$2:${CL('Fecha')}${FILAS}"
RANGO_T = f"Tickets!${TL('Fecha')}$2:${TL('Fecha')}${FILAS}"
RANGO_L = f"Compras!${CL('Importe')}$2:${CL('Importe')}${FILAS}"
RANGO_N = f"Compras!${CL('IVA $')}$2:${CL('IVA $')}${FILAS}"
RANGO_F = f"Compras!${CL('Categoria')}$2:${CL('Categoria')}${FILAS}"
RANGO_O = f"Compras!${CL('Forma de pago')}$2:${CL('Forma de pago')}${FILAS}"
C6 = f"$B${R_TOT_INI}"
C7 = f"$B${R_TOT_INI + 1}"

titulo_seccion(wr, f"A{R_TOT_TIT}", "ESTADO DEL EJERCICIO")
fila_etiqueta(wr, R_TOT_INI, "Gasto total del ejercicio",
              f'=ROUND(SUMIFS({RANGO_L},{RANGO_C},">="&{ANIO},{RANGO_C},"<="&{FIN_ANIO}),2)',
              co.MONEDA)
fila_etiqueta(wr, R_TOT_INI + 1, "Tickets registrados",
              f'=COUNTIFS({RANGO_T},">="&{ANIO},{RANGO_T},"<="&{FIN_ANIO})', "0")
fila_etiqueta(wr, R_TOT_INI + 2, "Lineas registradas",
              f'=COUNTIFS({RANGO_C},">="&{ANIO},{RANGO_C},"<="&{FIN_ANIO})', "0")
fila_etiqueta(wr, R_TOT_INI + 3, "Importe medio por ticket",
              f'=IF({C7}=0,0,ROUND({C6}/{C7},2))', co.MONEDA)
fila_etiqueta(wr, R_TOT_INI + 4, "IVA incluido en las compras",
              f'=ROUND(SUMIFS({RANGO_N},{RANGO_C},">="&{ANIO},{RANGO_C},"<="&{FIN_ANIO}),2)',
              co.MONEDA)
fila_etiqueta(wr, R_TOT_INI + 5, "Tickets a revisar (no cuadran)",
              f'=COUNTIFS({RANGO_T},">="&{ANIO},{RANGO_T},"<="&{FIN_ANIO},'
              f'Tickets!${TL("Estado")}$2:${TL("Estado")}${FILAS},"REVISAR")', "0")
fila_etiqueta(wr, R_TOT_INI + 6, "Posibles duplicados",
              f'=COUNTIFS({RANGO_T},">="&{ANIO},{RANGO_T},"<="&{FIN_ANIO},'
              f'Tickets!${TL("Duplicado")}$2:${TL("Duplicado")}${FILAS},"DUPLICADO")', "0")
for r in (R_TOT_INI + 5, R_TOT_INI + 6):
    wr.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f"B{r}>0"], fill=fill_rojo,
                                                       font=Font(color=ROJO_TXT, bold=True)))
    wr.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f"B{r}=0"], fill=fill_verde,
                                                       font=Font(color=VERDE_TXT)))

titulo_seccion(wr, f"A{R_MES_TIT}", "GASTO POR MES")
cabecera(wr, R_MES_CAB, ["Mes", "N. tickets", "Gasto", "% del ejercicio"])
for i, mes in enumerate(MESES):
    r = R_MES_INI + i
    ini = f'DATE($B$2,{i + 1},1)'
    wr.cell(row=r, column=1, value=mes)
    wr.cell(row=r, column=2, value=(
        f'=COUNTIFS({RANGO_T},">="&{ini},{RANGO_T},"<="&EOMONTH({ini},0))')).number_format = "0"
    wr.cell(row=r, column=3, value=(
        f'=ROUND(SUMIFS({RANGO_L},{RANGO_C},">="&{ini},{RANGO_C},"<="&EOMONTH({ini},0)),2)'
    )).number_format = co.MONEDA
    wr.cell(row=r, column=4, value=f'=IF({C6}=0,0,$C{r}/{C6})').number_format = co.PCT
    for col in range(1, 5):
        wr.cell(row=r, column=col).border = borde
wr.cell(row=R_MES_TOT, column=1, value="TOTAL").font = f_bold
for col in (2, 3, 4):
    c = wr.cell(row=R_MES_TOT, column=col,
                value=f"=SUM({get_column_letter(col)}{R_MES_INI}:{get_column_letter(col)}{R_MES_FIN})")
    c.font = f_bold
    c.fill = fill_total
    c.number_format = "0" if col == 2 else (co.MONEDA if col == 3 else co.PCT)

titulo_seccion(wr, f"A{R_CAT_TIT}", "GASTO POR CATEGORIA")
cabecera(wr, R_CAT_CAB, ["Categoria", "Gasto", "% del ejercicio", "N. lineas"])
for i in range(co.CAP_CAT):
    r = R_CAT_INI + i
    wr.cell(row=r, column=1, value=f'=IF(Listas!$A{i + 2}="","",Listas!$A{i + 2})')
    wr.cell(row=r, column=2, value=(
        f'=IF($A{r}="","",ROUND(SUMIFS({RANGO_L},{RANGO_F},$A{r},'
        f'{RANGO_C},">="&{ANIO},{RANGO_C},"<="&{FIN_ANIO}),2))')).number_format = co.MONEDA
    wr.cell(row=r, column=3, value=f'=IF(OR($A{r}="",{C6}=0),"",$B{r}/{C6})').number_format = co.PCT
    wr.cell(row=r, column=4, value=(
        f'=IF($A{r}="","",COUNTIFS({RANGO_F},$A{r},'
        f'{RANGO_C},">="&{ANIO},{RANGO_C},"<="&{FIN_ANIO}))')).number_format = "0"
    for col in range(1, 5):
        wr.cell(row=r, column=col).border = borde
wr.cell(row=R_CAT_SIN, column=1, value="Sin clasificar").font = f_bold
wr.cell(row=R_CAT_SIN, column=2, value=(
    f'=ROUND(SUMIFS({RANGO_L},{RANGO_C},">="&{ANIO},{RANGO_C},"<="&{FIN_ANIO},'
    f'{RANGO_F},""),2)')).number_format = co.MONEDA
wr.cell(row=R_CAT_SIN, column=3, value=f'=IF({C6}=0,0,$B{R_CAT_SIN}/{C6})').number_format = co.PCT
wr.cell(row=R_CAT_SIN, column=4, value=(
    f'=COUNTIFS({RANGO_C},">="&{ANIO},{RANGO_C},"<="&{FIN_ANIO},{RANGO_F},"")')).number_format = "0"
wr.cell(row=R_CAT_TOT, column=1, value="TOTAL").font = f_bold
for col in (2, 3, 4):
    c = wr.cell(row=R_CAT_TOT, column=col,
                value=f"=SUM({get_column_letter(col)}{R_CAT_INI}:{get_column_letter(col)}{R_CAT_SIN})")
    c.font = f_bold
    c.fill = fill_total
    c.number_format = "0" if col == 4 else (co.MONEDA if col == 2 else co.PCT)

titulo_seccion(wr, f"A{R_PAG_TIT}", "GASTO POR FORMA DE PAGO")
cabecera(wr, R_PAG_CAB, ["Forma de pago", "Gasto", "% del ejercicio"])
for i in range(co.CAP_PAGO):
    r = R_PAG_INI + i
    wr.cell(row=r, column=1, value=f'=IF(Listas!$B{i + 2}="","",Listas!$B{i + 2})')
    wr.cell(row=r, column=2, value=(
        f'=IF($A{r}="","",ROUND(SUMIFS({RANGO_L},{RANGO_O},$A{r},'
        f'{RANGO_C},">="&{ANIO},{RANGO_C},"<="&{FIN_ANIO}),2))')).number_format = co.MONEDA
    wr.cell(row=r, column=3, value=f'=IF(OR($A{r}="",{C6}=0),"",$B{r}/{C6})').number_format = co.PCT
    for col in range(1, 4):
        wr.cell(row=r, column=col).border = borde
wr.cell(row=R_PAG_TOT, column=1, value="TOTAL").font = f_bold
c = wr.cell(row=R_PAG_TOT, column=2,
            value=f"=SUM(B{R_PAG_INI}:B{R_PAG_FIN})")
c.font = f_bold
c.fill = fill_total
c.number_format = co.MONEDA

wr.cell(row=R_PAG_TOT + 2, column=1,
        value="Celdas amarillas = parametros editables (ejercicio y moneda).").font = f_nota

# --- graficos (flotan a la derecha de las tablas)
g_mes = BarChart()
g_mes.type = "col"
g_mes.title = "Gasto por mes"
g_mes.height, g_mes.width = 8.5, 15.5
g_mes.add_data(Reference(wr, min_col=3, min_row=R_MES_CAB, max_row=R_MES_FIN),
               titles_from_data=True)
g_mes.set_categories(Reference(wr, min_col=1, min_row=R_MES_INI, max_row=R_MES_FIN))
g_mes.legend = None
g_mes.y_axis.numFmt = '#,##0'
wr.add_chart(g_mes, "F3")

g_cat = PieChart()
g_cat.title = "Reparto por categoria"
g_cat.height, g_cat.width = 10.0, 15.5
g_cat.add_data(Reference(wr, min_col=2, min_row=R_CAT_CAB, max_row=R_CAT_FIN),
               titles_from_data=True)
g_cat.set_categories(Reference(wr, min_col=1, min_row=R_CAT_INI, max_row=R_CAT_FIN))
g_cat.legend = None
g_cat.dataLabels = DataLabelList()
g_cat.dataLabels.showCatName = True
g_cat.dataLabels.showPercent = True
wr.add_chart(g_cat, "F21")

# ---------------------------------------------------------------- Presupuesto
wp = wb.create_sheet("Presupuesto")
wp.column_dimensions["A"].width = 32
for col in "BCDE":
    wp.column_dimensions[col].width = 19
wp["A1"] = "PRESUPUESTO MENSUAL POR CATEGORIA"
wp["A1"].font = f_titulo
wp["A2"] = "Escribe tu presupuesto mensual (amarillo). El gasto se calcula sobre el mes en curso."
wp["A2"].font = f_nota
cabecera(wp, 4, ["Categoria", "Presupuesto mensual", "Gastado mes actual", "% ejecutado",
                 "Disponible"], [32, 19, 19, 19, 19])
MES_ACT = ('EOMONTH(TODAY(),-1)+1', 'EOMONTH(TODAY(),0)')
for i in range(co.CAP_CAT):
    r = 5 + i
    wp.cell(row=r, column=1, value=f'=IF(Listas!$A{i + 2}="","",Listas!$A{i + 2})')
    wp.cell(row=r, column=2).fill = fill_entrada
    wp.cell(row=r, column=2).number_format = co.MONEDA
    wp.cell(row=r, column=3, value=(
        f'=IF($A{r}="","",ROUND(SUMIFS({RANGO_L},{RANGO_F},$A{r},'
        f'{RANGO_C},">="&{MES_ACT[0]},{RANGO_C},"<="&{MES_ACT[1]}),2))')).number_format = co.MONEDA
    wp.cell(row=r, column=4,
            value=f'=IF(OR($A{r}="",N($B{r})=0),"",$C{r}/$B{r})').number_format = co.PCT
    wp.cell(row=r, column=5,
            value=f'=IF($A{r}="","",ROUND(N($B{r})-$C{r},2))').number_format = co.MONEDA
    for col in range(1, 6):
        wp.cell(row=r, column=col).border = borde
r_fin_pres = 5 + co.CAP_CAT - 1
wp.cell(row=r_fin_pres + 1, column=1, value="TOTAL").font = f_bold
for col in (2, 3, 5):
    c = wp.cell(row=r_fin_pres + 1, column=col,
                value=f"=SUM({get_column_letter(col)}5:{get_column_letter(col)}{r_fin_pres})")
    c.font = f_bold
    c.fill = fill_total
    c.number_format = co.MONEDA
wp.conditional_formatting.add(
    f"D5:D{r_fin_pres}",
    FormulaRule(formula=['AND($D5<>"",$D5>1)'], fill=fill_rojo, font=Font(color=ROJO_TXT, bold=True)))
wp.conditional_formatting.add(
    f"D5:D{r_fin_pres}",
    FormulaRule(formula=['AND($D5<>"",$D5<=1)'], fill=fill_verde, font=Font(color=VERDE_TXT)))

wb.move_sheet("Listas", offset=3)
wb.active = 0
os.makedirs(co.DIR_ENTRADA, exist_ok=True)
os.makedirs(co.DIR_ARCHIVADOS, exist_ok=True)
wb.save(XLSX)
print("OK ->", XLSX)
print("Hojas:", wb.sheetnames)
print(f"Resumen: totales {R_TOT_INI}-{R_TOT_FIN}, meses {R_MES_INI}-{R_MES_FIN}, "
      f"categorias {R_CAT_INI}-{R_CAT_FIN}, pagos {R_PAG_INI}-{R_PAG_FIN}")
