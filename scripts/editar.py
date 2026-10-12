# -*- coding: utf-8 -*-
"""Corrige o anula un ticket ya registrado.

Sirve para el caso real de haber leido mal un ticket (un importe, una linea,
una categoria) o para cuando una compra se devuelve. Nunca borra el escaneo
original: si se anula, se deja constancia en tickets/anulados.csv.

Uso:
  python editar.py --ver TK-0003
  python editar.py --anular TK-0003 --motivo "compra devuelta"
  python editar.py --reemplazar corregido.json --id TK-0003

El JSON de --reemplazar tiene el mismo formato que el del registro. Si no se
indica --id, se usa el campo "id" del propio JSON.

Opciones:
  --libro RUTA            libro a modificar
  --carpeta-tickets RUTA  raiz de los escaneos
  --sin-copia             no guarda copia de seguridad antes de escribir
  --seco                  simula sin escribir nada
"""
import argparse
import datetime as dt
import json
import os
import sys

from openpyxl import load_workbook

import comun as co
import registrar_ticket as reg

ENTRADA_COMPRAS = ["ID ticket", "Linea", "Fecha", "Comercio", "Categoria", "Descripcion",
                   "Cant.", "Unid.", "Precio unit.", "Descuento", "IVA %", "Forma de pago",
                   "Moneda", "N. ticket", "Archivo", "Notas"]
ENTRADA_TICKETS = ["ID ticket", "Fecha", "Comercio", "Total ticket", "N. ticket",
                   "Forma de pago", "Moneda", "Archivo escaneado", "Notas"]


class ErrorEdicion(Exception):
    pass


def localizar(wt, ws, idt):
    """Devuelve el ticket, su fila y las filas de sus lineas en Compras."""
    tickets = co.leer_tabla(wt, co.IT)
    cabecera = None
    for t in tickets:
        if co.normalizar(t["ID ticket"]).lower() == idt.lower():
            cabecera = t
            break
    if cabecera is None:
        disponibles = ", ".join(co.normalizar(t["ID ticket"]) for t in tickets) or "ninguno"
        raise ErrorEdicion(f"no existe el ticket {idt}. Registrados: {disponibles}")
    lineas = [l for l in co.leer_tabla(ws, co.IC)
              if co.normalizar(l["ID ticket"]).lower() == idt.lower()]
    if not lineas:
        raise ErrorEdicion(f"el ticket {idt} no tiene lineas en Compras")
    return cabecera, lineas


def limpiar(ws, fila, campos, indices):
    """Vacia las columnas de entrada y deja intactas las de formula.

    Ojo: ws.cell(..., value=None) NO borra la celda en openpyxl; hay que
    asignar el valor despues de obtenerla.
    """
    for campo in campos:
        ws.cell(row=fila, column=indices[campo] + 1).value = None


def describir(cabecera, lineas, titulo="TICKET"):
    print(f"{titulo} {co.normalizar(cabecera['ID ticket'])}")
    fecha = cabecera["Fecha"]
    fecha = fecha.date() if isinstance(fecha, dt.datetime) else fecha
    print(f"  Fecha:    {fecha}   Comercio: {co.normalizar(cabecera['Comercio'])}")
    print(f"  Ticket:   {co.normalizar(cabecera['N. ticket'])}   "
          f"Pago: {co.normalizar(cabecera['Forma de pago'])}   "
          f"Moneda: {co.normalizar(cabecera['Moneda'])}")
    print(f"  Escaneo:  {co.normalizar(cabecera['Archivo escaneado']) or '(ninguno)'}")
    suma = 0.0
    for l in lineas:
        imp = co.importe_linea(l["Cant."], l["Precio unit."], l["Descuento"])
        suma += imp
        cant = l["Cant."]
        cant_txt = f"{cant:g}" if isinstance(cant, (int, float)) else str(cant or "")
        print(f"    fila {l['_fila']:>4}  {co.normalizar(l['Descripcion'])[:34]:<34} "
              f"{cant_txt:>6} {co.normalizar(l['Unid.']):<4} {imp:>12,.2f}  "
              f"{co.normalizar(l['Categoria'])}")
    total = cabecera["Total ticket"] or 0
    print(f"  Total del papel: {total:>12,.2f}")
    print(f"  Suma de lineas:  {suma:>12,.2f}")
    print(f"  Diferencia:      {total - suma:>+12,.2f}  ->  "
          f"{'OK' if abs(total - suma) <= 0.02 else 'REVISAR'}")
    return suma


def claves_de(tickets):
    claves = {}
    for t in tickets:
        try:
            claves[co.clave_duplicado(t["Fecha"], t["Comercio"], t["Total ticket"],
                                      t["N. ticket"])] = co.normalizar(t["ID ticket"])
        except (ValueError, TypeError):
            continue
    return claves


def reemplazar(args, wb, ws, wt, wl, cabecera, lineas):
    idt = co.normalizar(cabecera["ID ticket"])
    if not os.path.exists(args.reemplazar):
        raise ErrorEdicion(f"no encuentro el archivo de datos: {args.reemplazar}")
    with open(args.reemplazar, encoding="utf-8") as f:
        datos = json.load(f)
    if isinstance(datos, list):
        if len(datos) != 1:
            raise ErrorEdicion("--reemplazar admite un solo ticket")
        datos = datos[0]
    datos.setdefault("id", idt)

    categorias = reg.leer_lista(wl, 1, co.CAP_CAT)
    pagos = reg.leer_lista(wl, 2, co.CAP_PAGO)
    unidades = reg.leer_lista(wl, 3, co.CAP_UNID)
    monedas = reg.leer_lista(wl, 4, 11)
    nuevo = reg.validar(datos, wl, categorias, pagos, unidades, monedas)
    nuevo["id"] = idt
    if not nuevo["archivo"]:
        # corregir un importe no debe perder el escaneo que ya estaba archivado
        nuevo["archivo"] = co.normalizar(cabecera["Archivo escaneado"]) or None

    # duplicados: se compara con los DEMAS tickets, no consigo mismo
    otros = [t for t in co.leer_tabla(wt, co.IT)
             if co.normalizar(t["ID ticket"]).lower() != idt.lower()]
    claves = claves_de(otros)
    try:
        clave = co.clave_duplicado(nuevo["fecha"], nuevo["comercio"], nuevo["total"],
                                   nuevo["numero"])
    except (ValueError, TypeError):
        clave = None
    if clave is not None and clave in claves and not args.permitir_duplicado:
        raise ErrorEdicion(
            f"el ticket corregido coincidiria con {claves[clave]} (misma fecha, comercio, "
            "total y numero). Si son compras distintas, repite con --permitir-duplicado."
        )

    print("--- antes ---")
    describir(cabecera, lineas)
    print("--- despues ---")

    anotaciones = []
    if nuevo["archivo"]:
        if args.seco:
            destino = reg.destino_previsto(nuevo["archivo"], nuevo["fecha"], nuevo["comercio"],
                                           nuevo["total"], args.carpeta_tickets)
        else:
            destino = reg.archivar(nuevo["archivo"], nuevo["fecha"], nuevo["comercio"],
                                   nuevo["total"], args.carpeta_tickets)
            anotaciones.append({
                "id": idt, "fecha": nuevo["fecha"].isoformat(),
                "archivo": co.ruta_legible(destino), "bytes": os.path.getsize(destino),
                "sha256": co.sha256(destino),
                "registrado_en": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")})
        nuevo["archivo"] = co.ruta_legible(destino)

    if not args.seco:
        for l in lineas:                      # liberar las lineas antiguas
            limpiar(ws, l["_fila"], ENTRADA_COMPRAS, co.IC)
        cursor = 2
        nuevo["filas_compras"] = []
        for _ in nuevo["lineas"]:
            r = co.fila_libre(ws, cursor)
            nuevo["filas_compras"].append(r)
            cursor = r + 1
        nuevo["fila_ticket"] = cabecera["_fila"]
        reg.escribir_compras(ws, nuevo)
        reg.escribir_ticket(wt, nuevo)
        if not args.sin_copia:
            args.copia = co.copia_seguridad(args.libro, "edicion")
        wb.save(args.libro)
        co.anotar_manifiesto(anotaciones, os.path.join(args.carpeta_tickets, "manifiesto.csv"))

    describe_nuevo(nuevo)
    return 0


def describe_nuevo(nuevo):
    print(f"  Total papel:  {nuevo['total']:>12,.2f}")
    print(f"  Total lineas: {nuevo['registrado']:>12,.2f}")
    print(f"  Diferencia:   {nuevo['diferencia']:>+12,.2f}  ->  "
          f"{'OK' if abs(nuevo['diferencia']) <= reg.AVISO_DIF else 'REVISAR'}")
    print(f"  Lineas nuevas en las filas "
          f"{', '.join(str(r) for r in nuevo['filas_compras'])} de Compras")


def anular(args, wb, ws, wt, cabecera, lineas):
    idt = co.normalizar(cabecera["ID ticket"])
    print("--- se anula ---")
    describir(cabecera, lineas)
    fecha = cabecera["Fecha"]
    fecha = fecha.date() if isinstance(fecha, dt.datetime) else fecha
    if not args.seco:
        for l in lineas:
            limpiar(ws, l["_fila"], ENTRADA_COMPRAS, co.IC)
        limpiar(wt, cabecera["_fila"], ENTRADA_TICKETS, co.IT)
        if not args.sin_copia:
            args.copia = co.copia_seguridad(args.libro, "anulacion")
        wb.save(args.libro)
        co.anotar_anulado({
            "id": idt, "fecha": fecha.isoformat() if fecha else "",
            "comercio": co.normalizar(cabecera["Comercio"]),
            "total": cabecera["Total ticket"],
            "archivo": co.normalizar(cabecera["Archivo escaneado"]),
            "motivo": args.motivo or "sin motivo indicado",
            "anulado_en": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }, os.path.join(args.carpeta_tickets, "anulados.csv"))
    print(f"  Anulado: se han quitado {len(lineas)} lineas y la fila del ticket.")
    print("  El escaneo NO se borra: sigue en la carpeta archivados.")
    return 0


def main():
    ap = argparse.ArgumentParser(description="Corrige o anula un ticket registrado.")
    ap.add_argument("--ver", metavar="ID", help="muestra un ticket y sus lineas")
    ap.add_argument("--reemplazar", metavar="JSON", help="corrige un ticket con datos nuevos")
    ap.add_argument("--anular", metavar="ID", help="quita un ticket del libro")
    ap.add_argument("--id", help="ticket a corregir (si el JSON no lo trae)")
    ap.add_argument("--motivo", default="", help="motivo de la anulacion")
    ap.add_argument("--libro", default=co.XLSX)
    ap.add_argument("--carpeta-tickets", default=co.DIR_TICKETS)
    ap.add_argument("--permitir-duplicado", action="store_true")
    ap.add_argument("--sin-copia", action="store_true")
    ap.add_argument("--seco", action="store_true")
    args = ap.parse_args()

    acciones = [x for x in (args.ver, args.reemplazar, args.anular) if x]
    if len(acciones) != 1:
        ap.error("indica exactamente una accion: --ver, --reemplazar o --anular")
    if not os.path.exists(args.libro):
        raise ErrorEdicion(f"no existe el libro: {args.libro}")

    wb = load_workbook(args.libro)
    ws, wt, wl = wb["Compras"], wb["Tickets"], wb["Listas"]
    args.copia = None

    if args.ver:
        cabecera, lineas = localizar(wt, ws, args.ver)
        describir(cabecera, lineas)
    elif args.anular:
        cabecera, lineas = localizar(wt, ws, args.anular)
        anular(args, wb, ws, wt, cabecera, lineas)
    else:
        idt = args.id
        if not idt:
            with open(args.reemplazar, encoding="utf-8") as f:
                contenido = json.load(f)
            if isinstance(contenido, list) and len(contenido) == 1:
                contenido = contenido[0]
            idt = contenido.get("id") if isinstance(contenido, dict) else None
        if not idt:
            raise ErrorEdicion("no se sabe que ticket corregir: usa --id o pon \"id\" en el JSON")
        cabecera, lineas = localizar(wt, ws, idt)
        reemplazar(args, wb, ws, wt, wl, cabecera, lineas)

    if args.copia:
        print(f"Copia de seguridad del libro anterior: {co.ruta_legible(args.copia)}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (ErrorEdicion, reg.ErrorTicket) as e:
        print("ERROR: " + str(e), file=sys.stderr)
        sys.exit(2)
    except json.JSONDecodeError as e:
        print(f"ERROR: el JSON no es valido ({e})", file=sys.stderr)
        sys.exit(2)
