# -*- coding: utf-8 -*-
"""Prueba de extremo a extremo del control de compras.

Copia la plantilla a una carpeta temporal y ejercita el circuito completo con
los mismos programas que se usan en produccion:
  registrar (suelto y en lote) -> validar -> detectar duplicados -> archivar
  el escaneo -> leer un PDF -> recuadrar
Despues recalcula con LibreOffice y comprueba los resultados del Resumen.

Uso: python prueba_logica.py
"""
import datetime as dt
import glob
import json
import os
import shutil
import subprocess
import sys

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

import comun as co
import registrar_ticket as reg

SCRIPTS = os.path.dirname(os.path.abspath(__file__))
REGISTRAR = os.path.join(SCRIPTS, "registrar_ticket.py")
LEER = os.path.join(SCRIPTS, "leer_escaneo.py")
VERIFICAR = os.path.join(SCRIPTS, "verificar.py")
EDITAR = os.path.join(SCRIPTS, "editar.py")
TMP = os.path.join(SCRIPTS, "_prueba")
LIBRO = os.path.join(TMP, "libro.xlsx")
CARPETA_TK = os.path.join(TMP, "tickets")
ENTRADA = os.path.join(CARPETA_TK, "entrada")
ARCHIVADOS = os.path.join(CARPETA_TK, "archivados")
LECTURA = os.path.join(TMP, "lectura")
SALIDA = os.path.join(TMP, "recalculado.xlsx")
JPG = b"\xff\xd8\xff\xe0 prueba-de-escaneo"

ANIO = dt.date.today().year
ANT = ANIO - 1

# ---------------------------------------------------------------- datos de prueba
T1 = {
    "fecha": f"{ANIO}-03-15", "comercio": "Supermercado Central", "numero": "A-1024",
    "forma_pago": "Tarjeta debito", "moneda": "ARS", "total": 8.00, "iva_global": 10,
    "notas": "compra semanal",
    "lineas": [
        {"descripcion": "Leche entera 1 L", "cantidad": 2, "unidad": "ud",
         "precio": 1.20, "categoria": "Supermercado"},
        {"descripcion": "Pan de molde", "cantidad": 1, "unidad": "ud",
         "precio": 2.10, "categoria": "Supermercado"},
        {"descripcion": "Yogur pack 6", "cantidad": 3, "unidad": "pack",
         "precio": 1.50, "descuento": 1.00, "categoria": "Supermercado"},
    ],
}
T2 = {  # el papel dice 6.00 y las lineas suman 5.00 -> debe quedar REVISAR
    "fecha": f"{ANIO}-03-20", "comercio": "Farmacia del Sol", "numero": "B-77",
    "forma_pago": "Efectivo", "moneda": "ARS", "total": 6.00,
    "lineas": [
        {"descripcion": "Ibuprofeno 600", "cantidad": 1, "unidad": "ud",
         "precio": 5.00, "iva": 21, "categoria": "Salud y farmacia"},
    ],
}
T3 = {  # identico a T1: debe rechazarse salvo --permitir-duplicado
    **{k: v for k, v in T1.items() if k != "lineas"},
    "lineas": [dict(l) for l in T1["lineas"]],
}
T4 = {  # categoria nueva, unidad kg, sin IVA informado
    "fecha": f"{ANIO}-04-05", "comercio": "Feria del barrio", "numero": "C-3",
    "forma_pago": "Transferencia", "moneda": "ARS", "total": 12.00,
    "lineas": [
        {"descripcion": "Tomate perita", "cantidad": 3, "unidad": "kg",
         "precio": 4.00, "categoria": "Feria y granja"},
    ],
}
T5 = {  # ejercicio anterior: no debe contaminar los totales del actual
    "fecha": f"{ANT}-12-31", "comercio": "Ferreteria Norte", "numero": "D-9",
    "forma_pago": "Tarjeta credito", "moneda": "ARS", "total": 100.00,
    "lineas": [
        {"descripcion": "Tornillos surtidos", "cantidad": 1, "unidad": "caja",
         "precio": 100.00, "categoria": "Hogar y ferreteria"},
    ],
}
T6 = {  # descuento impreso en renglon propio + IVA impreso que debe cuadrar
    "fecha": f"{ANIO}-03-25", "comercio": "Supermercado Central", "numero": "A-2000",
    "forma_pago": "Tarjeta debito", "moneda": "ARS", "total": 8.00,
    "iva_global": 21, "iva_importe": 1.39,
    "lineas": [
        {"descripcion": "Leche entera 1 L", "cantidad": 2, "unidad": "ud",
         "precio": 1.20, "categoria": "Supermercado"},
        {"descripcion": "Pan de molde", "cantidad": 1, "unidad": "ud",
         "precio": 2.10, "categoria": "Supermercado"},
        {"descripcion": "Yogur pack 6", "cantidad": 3, "unidad": "pack",
         "precio": 1.50, "categoria": "Supermercado"},
        {"descripcion": "Descuento general", "cantidad": 1, "unidad": "ud",
         "precio": -1.00, "categoria": "Supermercado"},
    ],
}
T7 = {  # dos tipos de IVA distintos: no se puede cotejar un unico tipo
    "fecha": f"{ANIO}-03-28", "comercio": "Almacen Don Pedro", "numero": "P-5",
    "forma_pago": "Efectivo", "moneda": "ARS", "total": 10.00,
    "lineas": [
        {"descripcion": "Vino tinto", "cantidad": 1, "unidad": "ud",
         "precio": 6.00, "iva": 21, "categoria": "Almacen y bebidas"},
        {"descripcion": "Galletitas", "cantidad": 1, "unidad": "ud",
         "precio": 4.00, "iva": 10.5, "categoria": "Almacen y bebidas"},
    ],
}
T9_MALO = {  # unidad inexistente: debe tumbar el lote entero
    "fecha": f"{ANIO}-04-10", "comercio": "Kiosco de la esquina", "numero": "K-1",
    "forma_pago": "Efectivo", "moneda": "ARS", "total": 3.00,
    "lineas": [
        {"descripcion": "Revista", "cantidad": 1, "unidad": "zzz",
         "precio": 3.00, "categoria": "Ocio y entretenimiento"},
    ],
}
T10 = {  # dos veces el mismo ticket dentro de un lote
    "fecha": f"{ANIO}-04-12", "comercio": "Panaderia La Espiga", "numero": "E-1",
    "forma_pago": "Efectivo", "moneda": "ARS", "total": 4.00,
    "lineas": [
        {"descripcion": "Docena de facturas", "cantidad": 1, "unidad": "ud",
         "precio": 4.00, "categoria": "Panaderia"},
    ],
}
T11 = {  # escaneo inexistente: el lote no debe mover nada
    "fecha": f"{ANIO}-04-15", "comercio": "Verduleria Verde", "numero": "F-1",
    "forma_pago": "Efectivo", "moneda": "ARS", "total": 6.00,
    "lineas": [
        {"descripcion": "Manzanas", "cantidad": 2, "unidad": "kg",
         "precio": 3.00, "categoria": "Verduleria y frutas"},
    ],
}
T12 = {  # acompana a T11 en el lote fallido
    "fecha": f"{ANIO}-04-16", "comercio": "Verduleria Verde", "numero": "F-2",
    "forma_pago": "Efectivo", "moneda": "ARS", "total": 5.00,
    "lineas": [
        {"descripcion": "Naranjas", "cantidad": 1, "unidad": "kg",
         "precio": 5.00, "categoria": "Verduleria y frutas"},
    ],
}

REGISTRADOS = [T1, T2, T3, T4, T5, T6, T7]          # los que deben quedar
CAT_NUEVA = "Feria y granja"
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

fallos = []
total_checks = 0


def comprobar(nombre, obtenido, esperado):
    global total_checks
    total_checks += 1
    if isinstance(esperado, float) and isinstance(obtenido, (int, float)):
        ok = abs(float(obtenido) - esperado) < 0.005
    else:
        ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(("OK    " if ok else "FALLA ") + f"| {nombre} -> {obtenido!r} (esperado {esperado!r})")


def imp(l):
    return round(l["cantidad"] * l["precio"] - l.get("descuento", 0), 2)


def iva_linea(t, l):
    pct = l.get("iva", t.get("iva_global"))
    if pct in (None, ""):
        return 0.0
    i = imp(l)
    return round(i - i / (1 + pct / 100), 2)


def del_anio(t):
    return dt.date.fromisoformat(t["fecha"]).year == ANIO


def registrar(datos, extra=()):
    ruta = os.path.join(TMP, "datos.json")
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False)
    cmd = [sys.executable, REGISTRAR, ruta, "--libro", LIBRO,
           "--carpeta-tickets", CARPETA_TK, *extra]
    return subprocess.run(cmd, capture_output=True, text=True)


def escaneo(nombre):
    ruta = os.path.join(ENTRADA, nombre)
    os.makedirs(ENTRADA, exist_ok=True)
    with open(ruta, "wb") as f:
        f.write(JPG)
    return ruta


def estado_libro():
    """(tickets, lineas) escritos ahora mismo en el libro."""
    wb = load_workbook(LIBRO)
    n_t = len(reg.leer_registrados(wb["Tickets"]))
    n_c = sum(1 for (v,) in wb["Compras"].iter_rows(min_row=2, max_row=co.FILAS, min_col=1,
                                                    max_col=1, values_only=True)
              if v not in (None, ""))
    wb.close()
    return n_t, n_c


def preparar():
    if os.path.exists(TMP):
        shutil.rmtree(TMP)
    os.makedirs(ENTRADA)
    wb = load_workbook(co.XLSX, read_only=True)
    for hoja in ("Compras", "Tickets"):
        if wb[hoja]["A2"].value not in (None, ""):
            wb.close()
            sys.exit(f"La plantilla {co.XLSX} ya tiene datos: la prueba necesita "
                     "el libro vacio.")
    wb.close()
    shutil.copyfile(co.XLSX, LIBRO)


def fila_de(ws, etiqueta, desde=1, hasta=250, col=1):
    for r in range(desde, hasta + 1):
        if co.normalizar(ws.cell(row=r, column=col).value) == etiqueta:
            return r
    raise AssertionError(f"no encontre la etiqueta {etiqueta!r} en {ws.title}")


def ultimo_total(ws, desde):
    filas = [r for r in range(desde, desde + 200)
             if co.normalizar(ws.cell(row=r, column=1).value) == "TOTAL"]
    if not filas:
        raise AssertionError("no encontre la fila TOTAL")
    return filas[-1]


def primer_total(ws, desde):
    for r in range(desde, desde + 200):
        if co.normalizar(ws.cell(row=r, column=1).value) == "TOTAL":
            return r
    raise AssertionError("no encontre la fila TOTAL")


# ================================================================ 1. REGISTRO
preparar()
print("== 1. registro individual, validacion, duplicados y archivado ==")

r = registrar({**T1, "archivo": escaneo("foto-1.jpg")})
comprobar("registrar T1 termina sin error", r.returncode, 0)
print(r.stdout.strip())

r = registrar({**T2, "archivo": escaneo("foto-2.jpg")})
comprobar("registrar T2 termina sin error", r.returncode, 0)
comprobar("T2 avisa REVISAR por diferencia de 1.00", "REVISAR" in r.stdout, True)

nombre_t1 = co.nombre_archivo(T1["fecha"], T1["comercio"], T1["total"], ".jpg")
comprobar("el escaneo se archiva con nombre normalizado",
          os.path.exists(os.path.join(ARCHIVADOS, nombre_t1)), True)
comprobar("la carpeta de entrada queda vacia", sorted(os.listdir(ENTRADA)), [])

r = registrar({**T3, "archivo": escaneo("foto-1.jpg")})
comprobar("T3 duplicado se rechaza", r.returncode, 2)
comprobar("el rechazo explica que es duplicado", "DUPLICADO" in r.stderr.upper(), True)
comprobar("el duplicado rechazado no archiva el escaneo",
          os.path.exists(os.path.join(ENTRADA, "foto-1.jpg")), True)

r = registrar({**T3, "archivo": os.path.join(ENTRADA, "foto-1.jpg")},
              ("--permitir-duplicado",))
comprobar("--permitir-duplicado: T3 se registra", r.returncode, 0)
comprobar("dos escaneos del mismo ticket no se pisan (sufijo -2)",
          os.path.exists(os.path.join(ARCHIVADOS, nombre_t1.replace(".jpg", "-2.jpg"))), True)

r = registrar({**T4, "archivo": escaneo("foto-3.jpg")})
comprobar("registrar T4 (categoria nueva) termina sin error", r.returncode, 0)
comprobar("T4 avisa de la categoria nueva", CAT_NUEVA in r.stdout, True)

r = registrar(T5)
comprobar("registrar T5 (ejercicio anterior) termina sin error", r.returncode, 0)

malo = {**T1, "fecha": f"{ANIO}-05-01", "numero": "Z-1", "forma_pago": "Bitcoin"}
r = registrar(malo)
comprobar("una forma de pago invalida se rechaza", r.returncode, 2)
comprobar("el error nombra la forma de pago", "forma de pago" in r.stderr, True)

r = registrar({**T1, "fecha": f"{ANIO}-05-02", "numero": "Z-2", "total": "ocho"})
comprobar("un total no numerico se rechaza", r.returncode, 2)

r = registrar({**T1, "fecha": f"{ANIO}-05-03", "numero": "Z-3", "lineas": []})
comprobar("un ticket sin lineas se rechaza", r.returncode, 2)

# ================================================================ 2. LOTE
print()
print("== 2. registro en lote (todo o nada) ==")
antes_t, antes_c = estado_libro()

r = registrar([{**T6, "archivo": escaneo("foto-6.jpg")},
               {**T7, "archivo": escaneo("foto-7.jpg")}])
comprobar("el lote de 2 tickets validos se registra", r.returncode, 0)
comprobar("el lote informa del total", "2 tickets, 6 lineas" in r.stdout, True)
despues_t, despues_c = estado_libro()
comprobar("el lote escribe los 2 tickets", despues_t - antes_t, 2)
comprobar("el lote escribe las 6 lineas", despues_c - antes_c, 6)
comprobar("los escaneos del lote quedan archivados", sorted(os.listdir(ENTRADA)), [])

antes_t, antes_c = estado_libro()
r = registrar([{**T10, "archivo": escaneo("foto-10.jpg")}, T9_MALO])
comprobar("un lote con un ticket invalido se rechaza entero", r.returncode, 2)
comprobar("el error indica que no se registro nada",
          "no se ha registrado nada" in r.stderr, True)
comprobar("nada del lote invalido se escribe", estado_libro(), (antes_t, antes_c))
comprobar("el escaneo del lote invalido no se mueve",
          os.path.exists(os.path.join(ENTRADA, "foto-10.jpg")), True)
os.remove(os.path.join(ENTRADA, "foto-10.jpg"))

antes_t, antes_c = estado_libro()
r = registrar([T10, dict(T10)])
comprobar("un lote con dos tickets identicos se rechaza", r.returncode, 2)
comprobar("el duplicado interno se explica", "mismo lote" in r.stderr, True)
comprobar("nada del lote duplicado se escribe", estado_libro(), (antes_t, antes_c))

antes_t, antes_c = estado_libro()
r = registrar([{**T12, "archivo": escaneo("foto-12.jpg")},
               {**T11, "archivo": os.path.join(ENTRADA, "no-existe.jpg")}])
comprobar("un lote con un escaneo inexistente se rechaza", r.returncode, 2)
comprobar("nada del lote se escribe", estado_libro(), (antes_t, antes_c))
comprobar("el escaneo valido del lote fallido sigue en entrada",
          os.path.exists(os.path.join(ENTRADA, "foto-12.jpg")), True)
os.remove(os.path.join(ENTRADA, "foto-12.jpg"))

r = registrar({**T6, "fecha": f"{ANIO}-05-04", "numero": "A-9999", "iva_importe": 5.00,
               "archivo": None}, ("--seco",))
comprobar("--seco no escribe nada", r.returncode, 0)
comprobar("--seco avisa del IVA que no cuadra", "IVA impreso" in r.stdout, True)
comprobar("--seco no deja rastro", estado_libro(), (despues_t, despues_c))

# ================================================================ 3. AUDITORIA
print()
print("== 3. copia de seguridad, manifiesto del archivo y auditoria ==")

copias = sorted(glob.glob(os.path.join(TMP, "copias", "*.xlsx")))
comprobar("se guarda una copia del libro por cada escritura", len(copias), 6)

ruta_manifiesto = os.path.join(CARPETA_TK, "manifiesto.csv")
manifiesto = co.leer_manifiesto(ruta_manifiesto)
archivados = sorted(os.listdir(ARCHIVADOS))
comprobar("el manifiesto anota un escaneo por ticket archivado", len(manifiesto), 6)
comprobar("los escaneos estan en archivados", len(archivados), 6)
entrada_t1 = [m for m in manifiesto if m["id"] == "TK-0001"][0]
comprobar("el manifiesto guarda la huella SHA-256 del original",
          entrada_t1["sha256"], co.sha256(entrada_t1["archivo"]))
comprobar("el manifiesto guarda el tamanio del original",
          int(entrada_t1["bytes"]), os.path.getsize(entrada_t1["archivo"]))


def auditar(extra=()):
    return subprocess.run([sys.executable, VERIFICAR, "--libro", LIBRO,
                           "--carpeta-tickets", CARPETA_TK, *extra],
                          capture_output=True, text=True)


gasto_esperado = round(sum(imp(l) for t in REGISTRADOS if del_anio(t)
                           for l in t["lineas"]), 2)
r = auditar(("--json",))
comprobar("la auditoria no encuentra problemas de integridad", r.returncode, 0)
inf = json.loads(r.stdout) if r.stdout.strip().startswith("{") else {}
comprobar("la auditoria calcula por su cuenta el mismo gasto", inf.get("gasto_total"),
          gasto_esperado)
comprobar("la auditoria ve el ticket que no cuadra",
          any("el papel dice 6.00" in a for a in inf.get("avisos", [])), True)
comprobar("la auditoria ve el duplicado", len(inf.get("duplicados", [])), 1)
comprobar("la auditoria cuenta los escaneos archivados", len(inf.get("problemas", [])), 0)

# --- un escaneo que desaparece
perdido = os.path.join(ARCHIVADOS, archivados[0])
os.rename(perdido, perdido + ".bak")
r = auditar()
comprobar("la auditoria detecta un escaneo que falta", r.returncode, 1)
comprobar("explica cual falta", "falta el escaneo" in r.stdout, True)
os.rename(perdido + ".bak", perdido)

# --- un escaneo alterado
objetivo = os.path.join(ARCHIVADOS, archivados[1])
with open(objetivo, "rb") as f:
    original = f.read()
with open(objetivo, "ab") as f:
    f.write(b"alterado")
r = auditar()
comprobar("la auditoria detecta un escaneo alterado", r.returncode, 1)
comprobar("lo explica como huella distinta", "ha cambiado" in r.stdout, True)
with open(objetivo, "wb") as f:
    f.write(original)
comprobar("al restaurarlo la auditoria vuelve a estar limpia", auditar().returncode, 0)

# --- un escaneo que ningun ticket reclama
huerfano = os.path.join(ARCHIVADOS, "2020-01-01_ticket-fantasma_1.00.jpg")
with open(huerfano, "wb") as f:
    f.write(JPG)
r = auditar()
comprobar("la auditoria avisa de un escaneo sin ticket",
          "no lo reclama ningun ticket" in r.stdout, True)
comprobar("un huerfano es un aviso, no un problema de integridad", r.returncode, 0)
os.remove(huerfano)

# ================================================================ 4. PDF
print()
print("== 4. lectura de un PDF escaneado ==")
pdf = os.path.join(ENTRADA, "ticket-escaneado.pdf")
fuente = None
for nombre in ("consola.ttf", "arial.ttf"):
    try:
        fuente = ImageFont.truetype(os.path.join(r"C:\Windows\Fonts", nombre), 22)
        break
    except OSError:
        continue
img = Image.new("RGB", (760, 240), "white")
ImageDraw.Draw(img).text((30, 40), "SUPERMERCADO CENTRAL   TOTAL $ 8,00", fill="black",
                         font=fuente)
img.save(pdf, "PDF", resolution=150)
r = subprocess.run([sys.executable, LEER, pdf, "--salida", LECTURA],
                   capture_output=True, text=True)
comprobar("leer_escaneo.py procesa el PDF", r.returncode, 0)
paginas = []
for raiz, _, ficheros in os.walk(LECTURA):
    paginas += [os.path.join(raiz, f) for f in ficheros if f.lower().endswith(".png")]
comprobar("el PDF se convierte en una imagen", len(paginas), 1)
if paginas:
    with Image.open(paginas[0]) as im:
        minimo, maximo = im.convert("L").getextrema()
        tamano = im.size
    comprobar("la imagen no esta en blanco", minimo < 128 and maximo > 200, True)
    comprobar("la imagen tiene tamano utilizable", min(tamano) >= 200, True)
os.remove(pdf)

# ================================================================ 5. RESULTADOS
print()
print("== 5. recalculo con LibreOffice y comprobacion de resultados ==")
if os.path.exists(SALIDA):
    os.remove(SALIDA)
r = subprocess.run([co.NODE, co.LIBREOFFICE_CLI, "recalculate", "--input", LIBRO,
                    "--output", SALIDA], capture_output=True, text=True)
if r.returncode != 0:
    sys.exit("Fallo la recalculada: " + r.stderr.strip())
wb = load_workbook(SALIDA, data_only=True)
wr, wt, wc, wl, wp = wb["Resumen"], wb["Tickets"], wb["Compras"], wb["Listas"], wb["Presupuesto"]

# --- Compras
filas = [(t, l) for t in REGISTRADOS for l in t["lineas"]]
n_lineas_hoja = sum(1 for (v,) in wc.iter_rows(min_row=2, max_row=co.FILAS, min_col=1,
                                               max_col=1, values_only=True)
                    if v not in (None, ""))
comprobar("Compras: una fila por producto", n_lineas_hoja, len(filas))
for i, (t, l) in enumerate(filas, start=2):
    comprobar(f"Compras L{i}: importe = cantidad x precio - descuento",
              wc.cell(row=i, column=co.IC["Importe"] + 1).value, imp(l))
comprobar("Compras D2: columna Mes", wc.cell(row=2, column=co.IC["Mes"] + 1).value,
          f"{ANIO}-03")
comprobar("Compras N2: IVA incluido al 10% (heredado del ticket)",
          wc.cell(row=2, column=co.IC["IVA $"] + 1).value, iva_linea(T1, T1["lineas"][0]))
comprobar("Compras N5: IVA incluido al 21% (propio de la linea)",
          wc.cell(row=5, column=co.IC["IVA $"] + 1).value, iva_linea(T2, T2["lineas"][0]))
comprobar("Compras N9: sin IVA informado queda vacio",
          wc.cell(row=9, column=co.IC["IVA $"] + 1).value in (None, ""), True)
fila_desc = 2 + len(T1["lineas"]) + len(T2["lineas"]) + len(T3["lineas"]) \
    + len(T4["lineas"]) + len(T5["lineas"]) + 3
comprobar("Compras: la linea de descuento queda en negativo",
          wc.cell(row=fila_desc, column=co.IC["Importe"] + 1).value, -1.00)

# --- Tickets
ids = [(r_, wt.cell(row=r_, column=1).value)
       for r_ in range(2, 40) if wt.cell(row=r_, column=1).value not in (None, "")]
comprobar("Tickets: una fila por ticket", len(ids), len(REGISTRADOS))
comprobar("Tickets: IDs correlativos", [v for _, v in ids],
          [f"TK-{n:04d}" for n in range(1, len(REGISTRADOS) + 1)])
for (r_, idt), t in zip(ids, REGISTRADOS):
    registrado = round(sum(imp(l) for l in t["lineas"]), 2)
    dif = round(t["total"] - registrado, 2)
    comprobar(f"Tickets {idt}: total registrado desde Compras",
              wt.cell(row=r_, column=co.IT["Total registrado"] + 1).value, registrado)
    comprobar(f"Tickets {idt}: diferencia con el papel",
              wt.cell(row=r_, column=co.IT["Diferencia"] + 1).value, dif)
    comprobar(f"Tickets {idt}: estado",
              wt.cell(row=r_, column=co.IT["Estado"] + 1).value,
              "OK" if abs(dif) <= 0.02 else "REVISAR")
    comprobar(f"Tickets {idt}: marca de duplicado",
              wt.cell(row=r_, column=co.IT["Duplicado"] + 1).value or "",
              "DUPLICADO" if t in (T1, T3) else "")

comprobar("Listas: la categoria nueva quedo anadida",
          CAT_NUEVA in [co.normalizar(wl.cell(row=r_, column=1).value)
                        for r_ in range(2, co.CAP_CAT + 5)], True)
comprobar("los graficos del Resumen sobreviven al recalculo", len(wr._charts), 2)
comprobar("el grafico de meses es de barras",
          wr._charts[0].__class__.__name__, "BarChart")
comprobar("el grafico de categorias es circular",
          wr._charts[1].__class__.__name__, "PieChart")

# --- Resumen
def valor(etiqueta, col=2):
    return wr.cell(row=fila_de(wr, etiqueta), column=col).value


total_anio = round(sum(imp(l) for t in REGISTRADOS if del_anio(t) for l in t["lineas"]), 2)
tickets_anio = len([t for t in REGISTRADOS if del_anio(t)])
lineas_anio = sum(len(t["lineas"]) for t in REGISTRADOS if del_anio(t))
iva_anio = round(sum(iva_linea(t, l) for t in REGISTRADOS if del_anio(t)
                     for l in t["lineas"]), 2)

comprobar("Resumen: gasto total del ejercicio", valor("Gasto total del ejercicio"), total_anio)
comprobar("Resumen: tickets registrados", valor("Tickets registrados"), tickets_anio)
comprobar("Resumen: lineas registradas", valor("Lineas registradas"), lineas_anio)
comprobar("Resumen: importe medio por ticket", valor("Importe medio por ticket"),
          round(total_anio / tickets_anio, 2))
comprobar("Resumen: IVA incluido en las compras", valor("IVA incluido en las compras"), iva_anio)
comprobar("Resumen: tickets a revisar", valor("Tickets a revisar (no cuadran)"), 1)
comprobar("Resumen: posibles duplicados", valor("Posibles duplicados"), 2)

meses = {}
for t in REGISTRADOS:
    if not del_anio(t):
        continue
    m = dt.date.fromisoformat(t["fecha"]).month
    meses[m] = round(meses.get(m, 0) + sum(imp(l) for l in t["lineas"]), 2)
for mes, monto in meses.items():
    comprobar(f"Resumen: gasto de {MESES[mes - 1]}", valor(MESES[mes - 1], col=3), monto)
comprobar("Resumen: total de meses = gasto del ejercicio",
          wr.cell(row=primer_total(wr, fila_de(wr, "GASTO POR MES")), column=3).value,
          total_anio)

categorias, pagos = {}, {}
for t in REGISTRADOS:
    if not del_anio(t):
        continue
    for l in t["lineas"]:
        categorias[l["categoria"]] = round(categorias.get(l["categoria"], 0) + imp(l), 2)
        pagos[t["forma_pago"]] = round(pagos.get(t["forma_pago"], 0) + imp(l), 2)
for cat, monto in categorias.items():
    comprobar(f"Resumen: categoria {cat}", valor(cat), monto)
comprobar("Resumen: sin clasificar (importe)", valor("Sin clasificar"), 0)
comprobar("Resumen: sin clasificar (lineas)", valor("Sin clasificar", col=4), 0)
comprobar("Resumen: total de categorias = gasto del ejercicio",
          wr.cell(row=fila_de(wr, "Sin clasificar") + 1, column=2).value, total_anio)
for pago, monto in pagos.items():
    comprobar(f"Resumen: pago {pago}", valor(pago), monto)
comprobar("Resumen: total de formas de pago = gasto del ejercicio",
          wr.cell(row=ultimo_total(wr, fila_de(wr, "GASTO POR FORMA DE PAGO")),
                  column=2).value, total_anio)

# --- Presupuesto
fila_tot_pres = ultimo_total(wp, 1)
comprobar("Presupuesto: gasto del mes en curso (sin compras este mes)",
          wp.cell(row=fila_tot_pres, column=3).value, 0)

# ================================================================ 6. EDICION
print()
print("== 6. correccion y anulacion de tickets ==")


def editar(extra):
    return subprocess.run([sys.executable, EDITAR, *extra, "--libro", LIBRO,
                           "--carpeta-tickets", CARPETA_TK], capture_output=True, text=True)


def escribir_json(datos, nombre="datos-edicion.json"):
    ruta = os.path.join(TMP, nombre)
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False)
    return ruta


r = editar(("--ver", "TK-0006"))
comprobar("--ver muestra el ticket", r.returncode, 0)
comprobar("--ver lista sus lineas", "Descuento general" in r.stdout, True)
comprobar("--ver informa del cuadre", "Diferencia:" in r.stdout, True)

r = editar(("--ver", "TK-9999"))
comprobar("--ver con un ID inexistente falla", r.returncode, 2)

# --- corregir el ticket que no cuadraba: el papel decia 6.00 y eran 5.00
antes_t, antes_c = estado_libro()
r = editar(("--reemplazar", escribir_json({**T2, "total": 5.00}), "--id", "TK-0002"))
comprobar("--reemplazar corrige el ticket", r.returncode, 0)
comprobar("la correccion no cambia el numero de tickets", estado_libro()[0], antes_t)
comprobar("la correccion no cambia el numero de lineas", estado_libro()[1], antes_c)

r = auditar(("--json",))
inf = json.loads(r.stdout)
comprobar("tras corregir, ya no queda ninguna diferencia",
          any("el papel dice 6.00" in a for a in inf.get("avisos", [])), False)
# quedan 2 avisos: el duplicado (T1/T3) y T5 registrado sin escaneo
comprobar("quedan el duplicado y el ticket sin escaneo", len(inf.get("avisos", [])), 2)
comprobar("sigue el aviso del duplicado",
          any("duplicado" in a.lower() for a in inf.get("avisos", [])), True)
comprobar("el gasto no cambia al corregir el total", inf["gasto_total"], gasto_esperado)
comprobar("el ticket corregido conserva su escaneo",
          co.normalizar(co.leer_tabla(
              load_workbook(LIBRO)["Tickets"], co.IT)[1]["Archivo escaneado"]) != "", True)

# --- corregir hacia datos que chocan con otro ticket
choque = {**T2, "fecha": T7["fecha"], "comercio": T7["comercio"],
          "numero": T7["numero"], "total": T7["total"]}
r = editar(("--reemplazar", escribir_json(choque, "choque.json"), "--id", "TK-0002"))
comprobar("corregir hacia un duplicado se rechaza", r.returncode, 2)
comprobar("lo explica", "coincidiria" in r.stderr, True)
comprobar("el rechazo no toca el libro", estado_libro(), (antes_t, antes_c))

# --- anular un ticket (compra devuelta)
r = editar(("--anular", "TK-0004", "--motivo", "compra devuelta"))
comprobar("--anular quita el ticket", r.returncode, 0)
comprobar("una fila menos en Tickets", estado_libro()[0], antes_t - 1)
comprobar("una linea menos en Compras", estado_libro()[1], antes_c - 1)
anulados = co.leer_anulados(os.path.join(CARPETA_TK, "anulados.csv"))
comprobar("la anulacion queda registrada", [a["id"] for a in anulados], ["TK-0004"])
comprobar("la anulacion guarda el motivo", anulados[0]["motivo"], "compra devuelta")
comprobar("el escaneo del ticket anulado no se borra", len(os.listdir(ARCHIVADOS)), 6)

r = auditar(("--json",))
inf = json.loads(r.stdout)
comprobar("tras anular, la auditoria sigue sin problemas de integridad", r.returncode, 0)
comprobar("el gasto baja en los 12.00 del ticket anulado",
          inf["gasto_total"], round(gasto_esperado - 12.00, 2))
comprobar("la auditoria informa de la anulacion", inf["anulados"], 1)
comprobar("el escaneo anulado no sale como huerfano",
          any("no lo reclama" in a for a in inf.get("avisos", [])), False)
comprobar("el ticket anulado ya no cuenta como duplicado posible", inf["duplicados"], [[
    "TK-0001", "TK-0003"]])

print()
print(f"{total_checks - len(fallos)}/{total_checks} comprobaciones correctas")
if fallos:
    print(f"{len(fallos)} FALLOS:")
    for f in fallos:
        print("  - " + f)
    sys.exit(1)
print("TODO CORRECTO")
shutil.rmtree(TMP, ignore_errors=True)