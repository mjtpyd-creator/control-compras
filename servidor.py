# -*- coding: utf-8 -*-
"""Servidor web del control de compras, para usarlo desde el telefono.

Sirve una web movil (en la carpeta web/) y una API JSON que reutiliza toda la
logica del sistema: resumen, listas, tickets y registro de compras con foto.

Se usa en la RED LOCAL (casa): enciendelo y abre la direccion que imprime desde
el telefono, estando en la misma WiFi que este ordenador.

Uso:
  python servidor.py                 # libro real, puerto 8080
  python servidor.py --puerto 8090
  python servidor.py --libro X.xlsx --carpeta-tickets Y --puerto 8090

Endpoints:
  GET  /api/resumen    resumen (gasto, por mes, por categoria, avisos)
  GET  /api/tickets    lista de tickets con sus lineas
  GET  /api/listas     categorias, formas de pago, unidades y monedas
  POST /api/foto       sube una foto (data URL) y la guarda en entrada
  POST /api/tickets    registra un ticket (mismo formato que el JSON del CLI)
"""
import argparse
import base64
import datetime as dt
import json
import os
import socket
import sys
import threading
import uuid
from collections import defaultdict
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

SCRIPT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scripts")
sys.path.insert(0, SCRIPT_DIR)

import comun as co            # noqa: E402
import registrar_ticket as reg  # noqa: E402
import verificar as vf          # noqa: E402
import prediccion as prd        # noqa: E402
from openpyxl import load_workbook  # noqa: E402

RAIZ = os.path.dirname(os.path.abspath(__file__))
WEB = os.path.join(RAIZ, "web")
LOCK = threading.Lock()
MIMES = {".html": "text/html; charset=utf-8", ".js": "application/javascript",
         ".css": "text/css; charset=utf-8", ".webmanifest": "application/manifest+json",
         ".png": "image/png", ".jpg": "image/jpeg", ".svg": "image/svg+xml",
         ".json": "application/json; charset=utf-8"}


def ip_lan():
    """IP de esta maquina en la red local (para el enlace del telefono)."""
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            s.connect(("8.8.8.8", 80))
            return s.getsockname()[0]
        finally:
            s.close()
    except OSError:
        return "127.0.0.1"


def fecha_str(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v or "")


# ------------------------------------------------------------------ datos

def resumen():
    with LOCK:
        return vf.auditar(co.XLSX, co.DIR_TICKETS, co.MANIFIESTO, None)


def listas():
    with LOCK:
        wb = load_workbook(co.XLSX)
        ws = wb["Listas"]

        def columna(i, capacidad):
            return [co.normalizar(ws.cell(row=r, column=i).value)
                    for r in range(2, capacidad + 2)
                    if ws.cell(row=r, column=i).value not in (None, "")]

        out = {"categorias": columna(1, co.CAP_CAT),
               "formas_pago": columna(2, co.CAP_PAGO),
               "unidades": columna(3, co.CAP_UNID),
               "monedas": columna(4, 11)}
        wb.close()
    return out


def tickets():
    with LOCK:
        wb = load_workbook(co.XLSX)
        ts = co.leer_tabla(wb["Tickets"], co.IT)
        cs = co.leer_tabla(wb["Compras"], co.IC)
        wb.close()

    por_id = defaultdict(list)
    for l in cs:
        por_id[co.normalizar(l["ID ticket"])].append(l)

    salida = []
    for t in ts:
        idt = co.normalizar(t["ID ticket"])
        lineas = por_id.get(idt, [])
        suma = round(sum(co.importe_linea(l["Cant."], l["Precio unit."], l["Descuento"])
                         for l in lineas), 2)
        total = t["Total ticket"] or 0
        dif = round(total - suma, 2)
        salida.append({
            "id": idt,
            "fecha": fecha_str(t["Fecha"]),
            "comercio": co.normalizar(t["Comercio"]),
            "total": total,
            "registrado": suma,
            "diferencia": dif,
            "estado": "OK" if abs(dif) <= 0.02 else "REVISAR",
            "numero": co.normalizar(t["N. ticket"]),
            "forma_pago": co.normalizar(t["Forma de pago"]),
            "notas": co.normalizar(t["Notas"]),
            "lineas": [{
                "descripcion": co.normalizar(l["Descripcion"]),
                "importe": co.importe_linea(l["Cant."], l["Precio unit."], l["Descuento"]),
                "categoria": co.normalizar(l["Categoria"]),
            } for l in lineas],
        })
    salida.sort(key=lambda x: x["fecha"], reverse=True)
    return salida


def ticket_limpio(tk):
    return {
        "id": tk["id"], "fecha": tk["fecha"].isoformat(), "comercio": tk["comercio"],
        "total": tk["total"], "registrado": tk["registrado"],
        "diferencia": tk["diferencia"], "forma_pago": tk["forma_pago"],
        "moneda": tk["moneda"], "numero": tk["numero"],
        "categorias_nuevas": tk["categorias_nuevas"], "duplicado_de": tk["duplicado_de"],
        "iva_aviso": tk["iva_aviso"], "archivo": tk["archivo"],
    }


def registrar_ticket(datos):
    lista = datos if isinstance(datos, list) else [datos]
    with LOCK:
        preparados, copia = reg.registrar(lista, co.XLSX, co.DIR_TICKETS)
    return {"ok": True, "tickets": [ticket_limpio(tk) for tk in preparados],
            "copia": co.ruta_legible(copia) if copia else None}


def extraer_ocr(ruta):
    """Intenta leer comercio y total de un ticket con Tesseract."""
    try:
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = co.TESSERACT_CMD
        from PIL import Image
        img = Image.open(ruta)
        # Zoom para mejorar lectura OCR
        w, h = img.size
        factor = max(2.0, 1500 / w) if w < 1500 else 1.5
        img = img.resize((int(w * factor), int(h * factor)), Image.LANCZOS)
        texto = None
        mejor_texto = ""
        # Intentar 4 rotaciones para encontrar la mejor lectura
        for angulo in [0, 90, 180, 270]:
            img_rot = img.rotate(angulo, expand=True) if angulo else img
            t = pytesseract.image_to_string(img_rot, lang="spa+eng", config="--psm 6")
            if len(t) > len(mejor_texto):
                mejor_texto = t
        texto = mejor_texto
        import re
        # Buscar total: números grandes cerca de TOTAL, TOTAL:, IMPORTE, etc.
        total = None
        # Buscar patrones típicos: "TOTAL 1234.56", "Total: $ 1.234,56"
        patrones_total = [
            r"(?:TOTAL|IMPORTE|TOTAL A PAGAR)[:\s]*\$?\s*([0-9]{1,3}(?:[.,]\s*[0-9]{3})*[.,][0-9]{2})",
            r"(?:TOTAL|IMPORTE)[:\s]*\$?\s*([0-9]+[.,][0-9]{2})",
            r"\$?\s*([0-9]{1,3}(?:[.,]\s*[0-9]{3})*[.,][0-9]{2})\s*(?:TOTAL|FIN|PAGO)",
            r"([0-9]+[.,][0-9]{2})",
        ]
        texto_lineas = texto.splitlines()
        # Buscar comercio: primera línea significativa o líneas cortas al inicio
        comercio = None
        for linea in texto_lineas[:10]:
            linea = linea.strip()
            if len(linea) > 2 and len(linea) < 80:
                # Evitar fechas, números de ticket, etc.
                if not re.match(r"^[0-9\/\-:.\s]+$", linea) and \
                   not linea.startswith("N°") and not linea.startswith("Fecha") and \
                   not linea.startswith("Ticket"):
                    comercio = linea
                    if len(comercio) > 0:
                        break
        # Buscar total con patrones
        for pat in patrones_total:
            matches = re.findall(pat, texto, re.IGNORECASE)
            for m in matches:
                m_limpio = m.replace(".", "").replace(",", ".").replace(" ", "")
                try:
                    v = float(m_limpio)
                    if v > 5:  # probable total
                        total = v
                        break
                except:
                    continue
            if total:
                break
        # Si no encuentra con patrones, buscar el número grande más cercano al final
        if total is None:
            numeros = re.findall(r"([0-9]{1,3}(?:[.,][0-9]{3})*[.,][0-9]{2})", texto)
            for n in numeros:
                v_str = n.replace(".", "").replace(",", ".")
                try:
                    v = float(v_str)
                    if v > 10:
                        total = v
                except:
                    continue
        # Intentar extraer líneas de producto del texto (más patrones)
        items = []
        for linea in texto_lineas:
            linea = linea.strip()
            # Ignorar líneas claramente no-producto
            if re.search(r"^(\s*TOTAL\b|\s*IMPORTE\b|\s*FECHA\b|\s*TICKET\b|\s*COMERCIO\b|\s*Caja\b)", linea, re.I):
                continue
            # Patrón 1: descripción + cantidad + precio (ej: "Leche 2 x 1.20")
            m = re.search(r"(.+?)\s+([0-9]+)\s*[xX]\s*([0-9]+[.,][0-9]{2})", linea)
            if m and len(m.group(1)) > 2 and len(m.group(1)) < 60:
                desc = m.group(1).strip()
                precio = float(m.group(3).replace(",", "."))
                items.append({"descripcion": desc[:50], "cantidad": int(m.group(2)), "precio": precio})
                continue
            # Patrón 2: descripción + precio al final (ej: "Leche entera 1.20")
            # Buscar precio de 2 decimales en la línea, excluir si es total
            m_precio = re.search(r"([0-9]+[.,][0-9]{2})\s*$", linea)
            if m_precio:
                desc = linea[:m_precio.start()].strip()
                if len(desc) > 2 and len(desc) < 60:
                    # Evitar líneas que solo son números o fechas
                    if not re.search(r"^[0-9\/\-:.,\s]+$", desc) and not re.search(r"(TOTAL|IMPORTE|FIN|PAGO)\b", desc, re.I):
                        precio = float(m_precio.group(1).replace(",", "."))
                        # Buscar cantidad si aparece al inicio o con x
                        cant_match = re.search(r"^([0-9]+)\s+[xX]?\s*", desc) or re.search(r"([0-9]+)\s+ud", desc, re.I)
                        cantidad = int(cant_match.group(1)) if cant_match else 1
                        # Evitar duplicados exactos de descripción
                        if not any(i["descripcion"] == desc[:50] for i in items):
                            items.append({"descripcion": desc[:50], "cantidad": cantidad, "precio": precio})
        # Limitar y ordenar
        seen = set()
        unique_items = []
        for it in items:
            if it["descripcion"] not in seen:
                seen.add(it["descripcion"])
                unique_items.append(it)
        items = unique_items[:12]
        return {"texto": texto, "comercio": comercio, "total": total, "items": items, "error_ocr": None}
    except Exception as e:
        import traceback
        return {"texto": "", "comercio": None, "total": None, "error_ocr": str(e), "trace": traceback.format_exc()}


def guardar_foto(datos):
    data_url = datos.get("data_url") or ""
    if not data_url.startswith("data:image/"):
        raise ValueError("la foto debe venir como data URL de imagen")
    tipo = data_url[5:].split(";", 1)[0].split("/", 1)[-1].lower()
    ext = {"jpeg": "jpg", "jpg": "jpg", "png": "png", "webp": "webp", "gif": "gif"}.get(tipo, "jpg")
    _, b64 = data_url.split(",", 1)
    try:
        contenido = base64.b64decode(b64)
    except Exception:
        raise ValueError("la foto no se pudo decodificar")
    if not contenido:
        raise ValueError("la foto esta vacia")
    os.makedirs(co.DIR_ENTRADA, exist_ok=True)
    nombre = f"foto-{uuid.uuid4().hex[:10]}.{ext}"
    ruta = os.path.join(co.DIR_ENTRADA, nombre)
    with open(ruta, "wb") as f:
        f.write(contenido)
    ocr = extraer_ocr(ruta)
    return {"ok": True, "path": co.ruta_legible(ruta),
            "ocr": ocr, "comercio_sugerido": ocr.get("comercio"), "total_sugerido": ocr.get("total")}


# ------------------------------------------------------------------ servidor

class Handler(BaseHTTPRequestHandler):
    server_version = "ControlCompras/1.0"

    def _send(self, code, payload, tipo="application/json; charset=utf-8"):
        if isinstance(payload, (dict, list)):
            cuerpo = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        else:
            cuerpo = payload if isinstance(payload, bytes) else str(payload).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(cuerpo)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(cuerpo)

    def _json_error(self, code, mensaje):
        self._send(code, {"ok": False, "error": str(mensaje)})

    def _leer_json(self):
        try:
            n = int(self.headers.get("Content-Length", 0))
        except ValueError:
            n = 0
        datos = self.rfile.read(n) if n else b""
        if not datos:
            return {}
        return json.loads(datos.decode("utf-8"))

    def do_GET(self):
        ruta = self.path.split("?", 1)[0]
        try:
            if ruta == "/" or ruta == "/index.html":
                self._servir_archivo(os.path.join(WEB, "index.html"))
            elif ruta.startswith("/api/resumen"):
                self._send(200, resumen())
            elif ruta.startswith("/api/listas"):
                self._send(200, listas())
            elif ruta.startswith("/api/tickets"):
                self._send(200, tickets())
            elif ruta.startswith("/api/sugerencias"):
                with LOCK:
                    self._send(200, prd.sugerencias(co.XLSX))
            elif ruta.startswith("/api/"):
                self._json_error(404, "endpoint desconocido")
            else:
                self._servir_archivo(os.path.join(WEB, ruta.lstrip("/")))
        except FileNotFoundError:
            self._json_error(404, "archivo no encontrado")
        except Exception as e:  # pragma: no cover
            self._json_error(500, f"error interno: {e}")

    def do_POST(self):
        ruta = self.path.split("?", 1)[0]
        try:
            if ruta == "/api/foto":
                self._send(200, guardar_foto(self._leer_json()))
            elif ruta == "/api/tickets":
                self._send(200, registrar_ticket(self._leer_json()))
            else:
                self._json_error(404, "endpoint desconocido")
        except reg.ErrorTicket as e:
            self._send(400, {"ok": False, "error": str(e)})
        except ValueError as e:
            self._send(400, {"ok": False, "error": str(e)})
        except json.JSONDecodeError:
            self._json_error(400, "el JSON recibido no es valido")
        except Exception as e:  # pragma: no cover
            self._json_error(500, f"error interno: {e}")

    def _servir_archivo(self, ruta):
        ruta = os.path.normpath(ruta)
        if not ruta.startswith(os.path.normpath(WEB)):
            raise FileNotFoundError(ruta)
        if not os.path.isfile(ruta):
            raise FileNotFoundError(ruta)
        with open(ruta, "rb") as f:
            contenido = f.read()
        ext = os.path.splitext(ruta)[1].lower()
        self._send(200, contenido, MIMES.get(ext, "application/octet-stream"))

    def log_message(self, fmt, *args):  # salida breve
        sys.stderr.write("  %s %s\n" % (self.address_string(), fmt % args))


def main():
    ap = argparse.ArgumentParser(description="Web movil del control de compras.")
    ap.add_argument("--libro", default=co.XLSX)
    ap.add_argument("--carpeta-tickets", default=co.DIR_TICKETS)
    ap.add_argument("--puerto", type=int, default=8080)
    args = ap.parse_args()

    # deja que la API use el libro y la carpeta indicados (util para pruebas)
    co.XLSX = args.libro
    co.DIR_TICKETS = args.carpeta_tickets
    co.DIR_ENTRADA = os.path.join(args.carpeta_tickets, "entrada")
    co.DIR_ARCHIVADOS = os.path.join(args.carpeta_tickets, "archivados")
    co.MANIFIESTO = os.path.join(args.carpeta_tickets, "manifiesto.csv")
    if not os.path.exists(args.libro):
        sys.exit(f"No existe el libro: {args.libro}")

    httpd = ThreadingHTTPServer(("0.0.0.0", args.puerto), Handler)
    lan = ip_lan()
    print("Control de compras - web para el movil", flush=True)
    print(f"  En este PC:  http://localhost:{args.puerto}", flush=True)
    print(f"  En el telefono (misma WiFi):  http://{lan}:{args.puerto}", flush=True)
    print(f"  Libro:       {co.ruta_legible(args.libro)}", flush=True)
    print("  Para detener: Ctrl+C", flush=True)
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\nDetenido.")


if __name__ == "__main__":
    main()
